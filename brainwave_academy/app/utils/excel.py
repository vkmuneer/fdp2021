"""Bulk student import: Excel template generation + parsing.

Only five fields are mandatory (admission_no, name, class, division,
parent_whatsapp) so office staff can fill a sheet quickly; everything else
is optional and defaults sensibly, matching the single-student add form.
"""
import io

import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

MANDATORY_FIELDS = ["admission_no", "name", "class", "division", "parent_whatsapp"]

FIELD_LABELS = {
    "admission_no": "admission_no*",
    "name": "name*",
    "class": "class*",
    "division": "division*",
    "parent_whatsapp": "parent_whatsapp*",
    "parent_name": "parent_name",
    "address": "address",
    "place": "place",
    "school_name": "school_name",
    "dob": "dob (YYYY-MM-DD)",
    "admission_date": "admission_date (YYYY-MM-DD)",
    "discount_amount": "discount_amount",
    "discount_reason": "discount_reason",
    "base_fee_override": "base_fee_override",
}

ALL_FIELDS = list(FIELD_LABELS.keys())


def build_template(class_names):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Students"

    headers = [FIELD_LABELS[f] for f in ALL_FIELDS]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    ws.append(
        [
            "BW101", "Amal Krishna", class_names[0] if class_names else "9", "A",
            "9847012345", "Suresh Kumar", "Omassery", "Omassery",
            "Govt. Higher Secondary School, Omassery", "2011-05-14", "2026-06-01", 0, "", "",
        ]
    )
    for i, header in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(14, len(header) + 2)

    info = wb.create_sheet("Instructions")
    info.append(["Field", "Required?", "Notes"])
    for cell in info[1]:
        cell.font = Font(bold=True)
    info.append(["admission_no", "Yes", "Must be unique for every student"])
    info.append(["name", "Yes", "Student's full name"])
    info.append(["class", "Yes", f"One of: {', '.join(class_names)}"])
    info.append(["division", "Yes", "e.g. A, B - created automatically if it doesn't exist yet"])
    info.append(["parent_whatsapp", "Yes", "10-digit mobile number (country code optional)"])
    info.append(["parent_name", "No", "Parent / guardian name"])
    info.append(["address", "No", ""])
    info.append(["place", "No", "Village/town the student is from"])
    info.append(["school_name", "No", "The regular school the student studies in"])
    info.append(["dob", "No", "Format YYYY-MM-DD"])
    info.append(["admission_date", "No", "Format YYYY-MM-DD, defaults to today"])
    info.append(["discount_amount", "No", "Rupees, defaults to 0"])
    info.append(["discount_reason", "No", "e.g. Sibling discount"])
    info.append(["base_fee_override", "No", "Leave blank to use the class's base fee"])
    info.column_dimensions["A"].width = 20
    info.column_dimensions["B"].width = 12
    info.column_dimensions["C"].width = 55

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def _normalize_header(value):
    if value is None:
        return ""
    text = str(value).strip().lower()
    text = text.split("(")[0]  # drop "(YYYY-MM-DD)" hints
    text = text.replace("*", "").strip()
    text = text.replace(" ", "_")
    return text


def parse_upload(file_stream):
    """Parses an uploaded workbook into a list of row dicts with normalized
    keys (matching FIELD_LABELS keys) plus a "_row" spreadsheet row number.
    Blank rows are skipped. Raises ValueError on unreadable files."""
    try:
        wb = openpyxl.load_workbook(file_stream, data_only=True)
    except Exception as exc:  # noqa: BLE001 - surface as a friendly upload error
        raise ValueError(f"Not a valid Excel (.xlsx) file: {exc}") from exc

    ws = wb["Students"] if "Students" in wb.sheetnames else wb.worksheets[0]

    try:
        header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    except StopIteration:
        return []
    headers = [_normalize_header(h) for h in header_row]

    rows = []
    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row is None or all(cell in (None, "") for cell in row):
            continue
        record = {"_row": row_idx}
        for header, value in zip(headers, row):
            if header:
                record[header] = value
        rows.append(record)
    return rows
